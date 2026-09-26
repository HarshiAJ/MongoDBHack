"""Generate a synthetic, deliberately messy RMI (Raw Material Intelligence) dataset.

Models a fictional automotive Tier-1 supplier ("Veltra Automotive") whose purchasing
team consolidates seven sources: Sales, BOM, Purchasing, Market Intelligence, SAP BW,
Excel files and the RMI Tracker. Every company, contract and price here is fictional.

Known data-quality issues are injected on purpose and listed in data/GROUND_TRUTH.md so
the ingestion/validation agent can be evaluated against them.

Usage: python scripts/generate_sample_data.py
"""

import csv
import json
import math
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

SEED = 42
TODAY = date(2026, 9, 26)
START = date(2025, 1, 1)
ROOT = Path(__file__).resolve().parent.parent / "data"
RAW = ROOT / "raw"

rng = random.Random(SEED)

# ---------------------------------------------------------------------------
# Master data (the "truth" the messy sources are derived from)
# ---------------------------------------------------------------------------

# code, description, SAP MATNR, index, SAP material group
MATERIALS = [
    ("RM-AL-A380", "Aluminium alloy A380 ingot", "000000000040001001", "ALU", "RM-NF"),
    ("RM-AL-6061", "Aluminium 6061-T6 sheet 2.0mm", "000000000040001002", "ALU", "RM-NF"),
    ("RM-ST-HRC", "Steel hot-rolled coil S355", "000000000040002001", "HRC", "RM-FE"),
    ("RM-ST-CRC", "Steel cold-rolled coil DC04", "000000000040002002", "HRC", "RM-FE"),
    ("RM-CU-ROD", "Copper wire rod 8mm ETP", "000000000040003001", "CU", "RM-NF"),
    ("RM-ZN-SHG", "Zinc special high grade", "000000000040004001", "ZN", "RM-NF"),
    ("RM-PA66-GF30", "Polyamide 66 GF30 granulate", "000000000040005001", "PA66", "RM-POLY"),
    ("RM-NR-SMR20", "Natural rubber SMR20", "000000000040006001", "RUBBER", "RM-ELAST"),
]
MAT = {m[0]: m for m in MATERIALS}

# How each source spells the material code (crosswalk the agent must learn)
BUYER_ALIASES = {
    "RM-AL-A380": ["AL A380", "Alu A380", "A380 ingot"],
    "RM-AL-6061": ["AL 6061", "Alu sheet 6061", "6061-T6"],
    "RM-ST-HRC": ["HRC S355", "Steel HRC", "HR coil"],
    "RM-ST-CRC": ["CRC DC04", "Steel CRC", "CR coil"],
    "RM-CU-ROD": ["Cu rod", "Copper 8mm", "CU wire rod"],
    "RM-ZN-SHG": ["Zinc SHG", "Zn"],
    "RM-PA66-GF30": ["PA66 GF30", "Nylon 66 GF30", "PA66"],
    "RM-NR-SMR20": ["SMR20", "Nat. rubber", "NR SMR 20"],
}
TRACKER_NAMES = {
    "RM-AL-A380": "Alu A380 ingot",
    "RM-AL-6061": "Alu 6061 sheet",
    "RM-ST-HRC": "Steel HRC",
    "RM-ST-CRC": "Steel CRC",
    "RM-CU-ROD": "Copper rod",
    "RM-ZN-SHG": "Zinc",
    "RM-PA66-GF30": "PA66 GF30",
    "RM-NR-SMR20": "Rubber SMR20",
}

# Finished goods and multi-level BOM (qty in kg per parent unit)
PRODUCTS = {
    "FG-BC-200": "Brake caliper BC-200",
    "FG-WH-EV1": "EV high-voltage wiring harness",
    "FG-SF-310": "Front seat frame SF-310",
    "FG-BE-500": "Battery enclosure BE-500",
    "FG-RD-120": "Radiator RD-120",
}
BOM = [
    # parent, child, qty, uom, level
    ("FG-BC-200", "SA-BC-HOUSING", 1, "EA", 1),
    ("SA-BC-HOUSING", "RM-AL-A380", 2.1, "KG", 2),
    ("FG-BC-200", "SA-BC-BRACKET", 1, "EA", 1),
    ("SA-BC-BRACKET", "RM-ST-CRC", 0.35, "KG", 2),
    ("FG-BC-200", "RM-NR-SMR20", 0.05, "KG", 1),
    ("FG-WH-EV1", "RM-CU-ROD", 3.8, "KG", 1),
    ("FG-WH-EV1", "RM-PA66-GF30", 0.9, "KG", 1),
    ("FG-SF-310", "RM-ST-CRC", 6.5, "KG", 1),
    ("FG-SF-310", "RM-ST-HRC", 2.0, "KG", 1),
    ("FG-SF-310", "RM-ZN-SHG", 0.12, "KG", 1),
    ("FG-BE-500", "SA-BE-TRAY", 1, "EA", 1),
    ("SA-BE-TRAY", "RM-AL-6061", 18.0, "KG", 2),
    ("FG-BE-500", "RM-ST-HRC", 1.5, "KG", 1),
    ("FG-BE-500", "RM-NR-SMR20", 0.4, "KG", 1),
    ("FG-RD-120", "RM-AL-6061", 3.2, "KG", 1),
    ("FG-RD-120", "RM-CU-ROD", 0.4, "KG", 1),
    ("FG-RD-120", "RM-PA66-GF30", 0.6, "KG", 1),
]

CUSTOMERS = {
    "C-NORDWERK": {
        "name": "Nordwerk Fahrzeugbau AG",
        "aliases": ["NORDWERK FAHRZEUGBAU AG", "Nordwerk", "Nordwerk FZB"],
        "country": "DE",
        "products": {"FG-BC-200": 42000, "FG-SF-310": 30000, "FG-RD-120": 18000},
    },
    "C-BRIGHTLINE": {
        "name": "Brightline Motors Inc.",
        "aliases": ["Brightline Motors", "BRIGHTLINE MOTORS INC", "Brightline"],
        "country": "US",
        "products": {"FG-WH-EV1": 25000, "FG-BE-500": 9000, "FG-BC-200": 15000},
    },
    "C-SAKURA": {
        "name": "Sakura EV Corporation",
        "aliases": ["Sakura EV Corp.", "SAKURA EV", "Sakura"],
        "country": "JP",
        "products": {"FG-BE-500": 7000, "FG-WH-EV1": 12000},
    },
}

# SAP LIFNR, name, country, invoicing currency, materials supplied
VENDORS = {
    "V-ALUCAST": ("0000100101", "AluCast Iberia S.L.", "ES", "EUR", ["RM-AL-A380"]),
    "V-MIDWEST": ("0000100102", "Midwest Diecast Alloys LLC", "US", "USD", ["RM-AL-A380"]),
    "V-NORDIC": ("0000100103", "Nordic Rolled Aluminium AB", "SE", "EUR", ["RM-AL-6061"]),
    "V-BALTIC": ("0000100104", "Baltic Steel Works AS", "EE", "EUR", ["RM-ST-HRC", "RM-ST-CRC"]),
    "V-CUPRUM": ("0000100105", "Cuprum Wire Sp. z o.o.", "PL", "PLN", ["RM-CU-ROD"]),
    "V-POLYNOVA": ("0000100106", "PolyNova Resins GmbH", "DE", "EUR", ["RM-PA66-GF30"]),
    "V-KEYSTONE": ("0000100107", "Keystone Rubber Co.", "MY", "USD", ["RM-NR-SMR20"]),
    "V-ZENITH": ("0000100108", "Zenith Zinc GmbH", "AT", "EUR", ["RM-ZN-SHG"]),
}

# ---------------------------------------------------------------------------
# Market indices: random walks with a late-2026 shock on aluminium, copper, PA66
# ---------------------------------------------------------------------------

INDEX_SPEC = {
    # name: (start, daily vol, unit, shock_start, shock_total)
    "ALU": (2350.0, 0.009, "USD/t", date(2026, 7, 20), 0.16),
    "CU": (9100.0, 0.010, "USD/t", date(2026, 8, 20), 0.10),
    "ZN": (2650.0, 0.011, "USD/t", None, 0.0),
    "HRC": (640.0, 0.006, "EUR/t", date(2026, 7, 1), -0.04),
    "PA66": (3.10, 0.004, "EUR/kg", date(2026, 8, 1), 0.07),
    "RUBBER": (170.0, 0.008, "USc/kg", date(2026, 8, 15), 0.04),
    "EURUSD": (1.08, 0.003, "USD per EUR", None, 0.0),
    "EURPLN": (4.30, 0.002, "PLN per EUR", None, 0.0),
}


def business_days(start, end):
    d = start
    while d <= end:
        if d.weekday() < 5:
            yield d
        d += timedelta(days=1)


def build_indices():
    days = list(business_days(START, TODAY - timedelta(days=1)))
    series = {}
    for name, (start, vol, _unit, shock_start, shock) in INDEX_SPEC.items():
        level = start
        out = {}
        shock_days = [d for d in days if shock_start and d >= shock_start]
        target = start
        for d in days:
            if shock_days and d >= shock_start:
                progress = (shock_days.index(d) + 1) / len(shock_days)
                target = start * (1 + shock * progress)
                level = max(level, target) if shock > 0 else min(level, target)
            # mean-reverting noise around the (possibly shocked) target level
            level *= 1 + rng.gauss(0, vol) + 0.05 * (target - level) / target
            out[d] = level
        series[name] = out
    return days, series


def month_avg(series, year, month):
    vals = [v for d, v in series.items() if d.year == year and d.month == month]
    return sum(vals) / len(vals)


# ---------------------------------------------------------------------------
# Contracts
# ---------------------------------------------------------------------------

VENDOR_CONTRACTS = [
    {
        "id": "VC-2025-014", "vendor": "V-ALUCAST", "material": "RM-AL-A380",
        "signed": "2025-01-15", "expires": "2027-12-31", "currency": "EUR",
        "base_price": 2.85, "price_unit": "EUR/kg", "index": "ALU",
        "base_index_period": "2024-12", "base_index_value": 2350.0,
        "share": 1.0, "trigger_pct": 5.0, "review": "monthly", "averaging": "prior calendar month",
        "fx": "ECB EURUSD monthly average", "notice_days": 15, "cap_pct": None,
        "volume_t": 1300,
    },
    {
        "id": "VC-2025-022", "vendor": "V-MIDWEST", "material": "RM-AL-A380",
        "signed": "2025-03-01", "expires": "2026-12-31", "currency": "USD",
        "base_price": 3.05, "price_unit": "USD/kg", "index": "ALU",
        "base_index_period": "2025-02", "base_index_value": 2380.0,
        "share": 0.9, "trigger_pct": 3.0, "review": "monthly", "averaging": "prior calendar month",
        "fx": None, "notice_days": 10, "cap_pct": 12.0,
        "volume_t": 400,
    },
    {
        "id": "VC-2024-031", "vendor": "V-NORDIC", "material": "RM-AL-6061",
        "signed": "2024-11-20", "expires": "2027-06-30", "currency": "EUR",
        "base_price": 4.40, "price_unit": "EUR/kg", "index": "ALU",
        "base_index_period": "2024-10", "base_index_value": 2400.0,
        "share": 0.85, "trigger_pct": 5.0, "review": "quarterly", "averaging": "prior 3 months",
        "fx": "ECB EURUSD monthly average", "notice_days": 30, "cap_pct": 15.0,
        "volume_t": 2600,
    },
    {
        "id": "VC-2025-007", "vendor": "V-BALTIC", "material": "RM-ST-HRC,RM-ST-CRC",
        "signed": "2025-01-05", "expires": "2027-01-04", "currency": "EUR",
        "base_price": 690.0, "price_unit": "EUR/t (HRC); CRC = HRC + 95 EUR/t extra", "index": "HRC",
        "base_index_period": "2024-12", "base_index_value": 640.0,
        "share": 1.0, "trigger_pct": 5.0, "review": "quarterly", "averaging": "prior 3 months",
        "fx": None, "notice_days": 30, "cap_pct": None,
        "volume_t": 3400,
        "amendment": {"id": "VC-2025-007-A1", "effective": "2026-07-01", "trigger_pct": 3.0,
                      "review": "monthly"},
    },
    {
        "id": "VC-2025-018", "vendor": "V-CUPRUM", "material": "RM-CU-ROD",
        "signed": "2025-02-10", "expires": "2027-02-09", "currency": "EUR",
        "base_price": 9.20, "price_unit": "EUR/kg (invoiced in PLN at NBP rate)", "index": "CU",
        "base_index_period": "2025-01", "base_index_value": 9100.0,
        "share": 1.0, "trigger_pct": 4.0, "review": "monthly", "averaging": "prior calendar month",
        "fx": "ECB EURUSD monthly average", "notice_days": 10, "cap_pct": None,
        "volume_t": 1250,
    },
    {
        "id": "VC-2025-026", "vendor": "V-POLYNOVA", "material": "RM-PA66-GF30",
        "signed": "2025-04-01", "expires": "2027-03-31", "currency": "EUR",
        "base_price": 3.95, "price_unit": "EUR/kg", "index": "PA66",
        "base_index_period": "2025-03", "base_index_value": 3.10,
        "share": 0.7, "trigger_pct": 6.0, "review": "quarterly", "averaging": "prior 3 months",
        "fx": None, "notice_days": 45, "cap_pct": 10.0,
        "volume_t": 380,
    },
    {
        "id": "VC-2024-040", "vendor": "V-KEYSTONE", "material": "RM-NR-SMR20",
        "signed": "2024-12-01", "expires": "2026-12-31", "currency": "USD",
        "base_price": 2.10, "price_unit": "USD/kg", "index": None,
        "base_index_period": None, "base_index_value": None,
        "share": 0.0, "trigger_pct": None, "review": "fixed price, no indexation",
        "averaging": None, "fx": None, "notice_days": 90, "cap_pct": None,
        "volume_t": 60,
    },
    {
        "id": "VC-2025-011", "vendor": "V-ZENITH", "material": "RM-ZN-SHG",
        "signed": "2025-01-20", "expires": "2027-01-19", "currency": "EUR",
        "base_price": 3.30, "price_unit": "EUR/kg", "index": "ZN",
        "base_index_period": "2024-12", "base_index_value": 2650.0,
        "share": 1.0, "trigger_pct": 5.0, "review": "monthly", "averaging": "prior calendar month",
        "fx": "ECB EURUSD monthly average", "notice_days": 15, "cap_pct": None,
        "volume_t": 45,
    },
]

CUSTOMER_CONTRACTS = [
    {
        "id": "CC-NW-2025-01", "customer": "C-NORDWERK", "signed": "2025-01-01",
        "expires": "2027-12-31",
        "surcharge": [
            {"index": "ALU", "materials": "aluminium (ALU index)", "share": 1.0, "trigger_pct": 5.0,
             "averaging": "3-month average", "lag": "applied from the first day of the following quarter"},
            {"index": "HRC", "materials": "steel (HRC index)", "share": 1.0, "trigger_pct": 5.0,
             "averaging": "3-month average", "lag": "applied from the first day of the following quarter"},
        ],
        "excluded": "copper, polymers and elastomers are included in the piece price and are not adjusted",
        "notice": "Supplier shall submit surcharge claims with index evidence no later than 20 days before quarter start.",
    },
    {
        "id": "CC-BL-2024-07", "customer": "C-BRIGHTLINE", "signed": "2024-07-01",
        "expires": "2027-06-30",
        "surcharge": [
            {"index": "ALU", "materials": "aluminium (ALU index)", "share": 0.8, "trigger_pct": 5.0,
             "averaging": "prior quarter average", "lag": "with a lag of one quarter"},
            {"index": "CU", "materials": "copper (CU index)", "share": 0.8, "trigger_pct": 5.0,
             "averaging": "prior quarter average", "lag": "with a lag of one quarter"},
        ],
        "excluded": "resins (including PA66), rubber and steel are fixed for the contract term",
        "notice": "Claims must be filed within 30 days after quarter end or are forfeited.",
    },
    {
        "id": "CC-SK-2025-03", "customer": "C-SAKURA", "signed": "2025-03-15",
        "expires": "2028-03-14",
        "surcharge": [
            {"index": "ALU", "materials": "aluminium (ALU index)", "share": 1.0, "trigger_pct": 8.0,
             "averaging": "6-month average", "lag": "semi-annual (1 April / 1 October)",
             "cap": "cumulative adjustment capped at 10% per contract year"},
        ],
        "excluded": "all other raw materials are fixed; copper exposure is borne by the supplier",
        "notice": "Adjustment requests require 45 days written notice before the adjustment date.",
    },
]


def pdf(path, title, blocks):
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 8)]
    for kind, content in blocks:
        if kind == "h":
            story.append(Paragraph(content, styles["Heading3"]))
        elif kind == "p":
            story.append(Paragraph(content, styles["BodyText"]))
            story.append(Spacer(1, 4))
        elif kind == "table":
            story.append(Table(content, hAlign="LEFT"))
            story.append(Spacer(1, 6))
    SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)


def write_vendor_contracts():
    out = RAW / "contracts" / "vendor"
    for c in VENDOR_CONTRACTS:
        lifnr, vname, country, _cur, _mats = VENDORS[c["vendor"]]
        mats = ", ".join(MAT[m][1] for m in c["material"].split(","))
        blocks = [
            ("p", f"<b>Contract no.</b> {c['id']} &nbsp; <b>Supplier no.</b> {lifnr}"),
            ("p", f"This Supply Agreement is made on {c['signed']} between <b>Veltra Automotive GmbH</b> "
                  f"(\"Buyer\") and <b>{vname}</b>, {country} (\"Supplier\")."),
            ("h", "1. Scope"),
            ("p", f"Supplier shall deliver {mats} to Buyer's plants in Wolfsried (DE) and Greenville (US). "
                  f"Estimated annual volume: {c['volume_t']:,} metric tonnes. Volumes are non-binding forecasts; "
                  f"Buyer commits to a minimum offtake of {int(c['volume_t'] * 0.7):,} t per contract year."),
            ("h", "2. Term"),
            ("p", f"The Agreement is effective from signature and expires on {c['expires']}. Either party may "
                  f"terminate with {c['notice_days'] * 2} days written notice. Renewal negotiations shall start "
                  f"no later than {c['notice_days']} days before expiry."),
            ("h", "3. Price"),
            ("p", f"The base price is {c['base_price']:.2f} {c['price_unit']}, delivered DAP Buyer's plant."),
        ]
        if c["index"]:
            cap = (f" The cumulative adjustment in any contract year shall not exceed {c['cap_pct']:.0f}% of the "
                   f"base price.") if c["cap_pct"] else ""
            fx = (" The relative index change is calculated in the index currency; no currency "
                  "conversion is applied to the adjustment.") if c["fx"] else ""
            blocks += [
                ("h", "4. Raw material price adjustment"),
                ("p", f"The price is linked to the {index_label(c['index'])}. The reference value is the "
                      f"{c['base_index_period']} monthly average of {c['base_index_value']:,} {INDEX_SPEC[c['index']][2]}. "
                      f"Adjustments are calculated {c['review']} using the index average over the {c['averaging']}."),
                ("p", f"An adjustment is triggered only when the index average deviates from the reference value "
                      f"by more than <b>{c['trigger_pct']:.0f}%</b> (up or down). {int(c['share'] * 100)}% of the "
                      f"relative index change is passed through to the base price.{cap}{fx}"),
                ("p", f"Supplier shall notify Buyer of any adjustment in writing at least {c['notice_days']} days "
                      f"before it takes effect, including index evidence. Adjustments without timely notice apply "
                      f"from the next review period."),
            ]
        else:
            blocks += [
                ("h", "4. Price firmness"),
                ("p", "Prices are fixed for the full term and are not subject to raw material, energy or "
                      "currency adjustment. Any new price after expiry must be agreed in writing."),
            ]
        blocks += [
            ("h", "5. Payment"),
            ("p", f"Payment 60 days end of month. Invoicing currency: "
                  f"{'PLN at the NBP mid-rate of the invoice date' if c['vendor'] == 'V-CUPRUM' else c['currency']}."),
            ("h", "6. Force majeure"),
            ("p", "Neither party is liable for delays caused by events beyond its reasonable control. "
                  "Market price movements do not constitute force majeure."),
        ]
        pdf(out / f"{c['id']}_{vname.split()[0]}.pdf", f"Supply Agreement {c['id']}", blocks)

        if "amendment" in c:
            a = c["amendment"]
            pdf(out / f"{a['id']}_{vname.split()[0]}_amendment.pdf", f"Amendment {a['id']} to {c['id']}", [
                ("p", f"Between Veltra Automotive GmbH and {vname}."),
                ("h", "1. Change to Section 4"),
                ("p", f"With effect from <b>{a['effective']}</b>, the adjustment trigger in Section 4 is reduced "
                      f"from {c['trigger_pct']:.0f}% to <b>{a['trigger_pct']:.0f}%</b> and the review frequency "
                      f"changes from {c['review']} to <b>{a['review']}</b>, using the prior calendar month average."),
                ("h", "2. Other terms"),
                ("p", "All other terms of the Agreement remain unchanged."),
            ])


def index_label(code):
    return {
        "ALU": "LME Aluminium cash settlement index (USD/t)",
        "CU": "LME Copper grade A cash settlement index (USD/t)",
        "ZN": "LME Zinc SHG cash settlement index (USD/t)",
        "HRC": "Northwest Europe HRC steel index (EUR/t)",
        "PA66": "European PA66 contract price assessment (EUR/kg)",
        "RUBBER": "SGX SMR20 rubber futures (USc/kg)",
    }[code]


def write_customer_contracts():
    out = RAW / "contracts" / "customer"
    for c in CUSTOMER_CONTRACTS:
        cust = CUSTOMERS[c["customer"]]
        parts = ", ".join(f"{PRODUCTS[p]} ({p})" for p in cust["products"])
        blocks = [
            ("p", f"<b>Agreement no.</b> {c['id']}"),
            ("p", f"Long-term supply agreement dated {c['signed']} between <b>{cust['name']}</b> (\"Customer\") "
                  f"and <b>Veltra Automotive GmbH</b> (\"Supplier\")."),
            ("h", "1. Parts"),
            ("p", f"Supplier supplies the following parts: {parts}. Piece prices are listed in Annex A."),
            ("h", "2. Term"),
            ("p", f"Valid until {c['expires']}."),
            ("h", "3. Raw material surcharge"),
        ]
        for s in c["surcharge"]:
            cap = f" {s['cap'].capitalize()}." if "cap" in s else ""
            blocks.append(("p", f"For {s['materials']}: if the {s['averaging']} of the index deviates by more than "
                                f"{s['trigger_pct']:.0f}% from the reference value in Annex B, "
                                f"{int(s['share'] * 100)}% of the resulting material cost change is passed through, "
                                f"{s['lag']}.{cap}"))
        blocks += [
            ("p", f"<b>Exclusions.</b> {c['excluded'][0].upper() + c['excluded'][1:]}."),
            ("h", "Annex B. Index reference values"),
            ("table", [["Index", "Reference period", "Reference value"]] + [
                [index_label(x["index"]), "2024-12 monthly average", f"{INDEX_SPEC[x['index']][0]:,}"]
                for x in c["surcharge"]]),
            ("h", "4. Claims procedure"),
            ("p", c["notice"]),
            ("h", "5. Annual price-down"),
            ("p", "Supplier grants a productivity price reduction of 2% on piece prices on each 1 January."),
        ]
        pdf(out / f"{c['id']}_{cust['name'].split()[0]}.pdf", f"Supply Agreement {c['id']}", blocks)


EMAIL = {}


def write_vendor_notice_email(idx):
    """Real-time change arriving as an unstructured document. The claimed price is deliberately
    higher than the contract formula allows, so the agent has something to catch."""
    contract = next(c for c in VENDOR_CONTRACTS if c["id"] == "VC-2025-014")
    correct = contract_price_kg(contract, "RM-AL-A380", date(2026, 10, 1), idx)
    claimed = round(correct * 1.045, 2)
    EMAIL.update(correct=correct, claimed=claimed)
    path = RAW / "contracts" / "inbox" / "2026-09-22_AluCast_price_notice.eml"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "From: key.accounts@alucast-iberia.example\n"
        "To: purchasing@veltra-automotive.example\n"
        "Date: Tue, 22 Sep 2026 09:14:00 +0200\n"
        "Subject: Price adjustment notice - A380 ingot - contract VC-2025-014\n"
        "\n"
        "Dear Purchasing team,\n\n"
        "Due to the sharp increase of the LME aluminium price since late July, we hereby notify you\n"
        "of a price adjustment according to section 4 of our agreement VC-2025-014.\n\n"
        "The September average to date exceeds the reference value by far more than the agreed threshold.\n"
        f"The new price of {claimed:.2f} EUR/kg applies to all deliveries from 1 October 2026.\n\n"
        "Please confirm receipt. Index evidence is attached.\n\n"
        "Kind regards,\nMarta Ruiz\nKey Account Manager, AluCast Iberia S.L.\n"
    )


# ---------------------------------------------------------------------------
# Tabular sources
# ---------------------------------------------------------------------------


def write_csv(path, header, rows, delimiter=","):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow(header)
        w.writerows(rows)


def bold_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")


def write_market(days, idx):
    out = RAW / "market_intelligence"
    # LME metals: ISO dates, long format, USD/t
    rows = [(d.isoformat(), metal, round(idx[code][d], 2))
            for d in days for code, metal in (("ALU", "Aluminium"), ("CU", "Copper"), ("ZN", "Zinc"))]
    # injected: one missing day, one duplicated day with a conflicting value
    rows = [r for r in rows if not (r[0] == "2026-03-17" and r[1] == "Copper")]
    dup = next(r for r in rows if r[0] == "2026-06-10" and r[1] == "Aluminium")
    rows.append((dup[0], dup[1], round(dup[2] * 1.018, 2)))
    write_csv(out / "lme_metals_daily.csv", ["date", "metal", "cash_usd_per_t"], rows)

    # Steel HRC: weekly Excel, dd/mm/yyyy strings, EUR/t
    wb = Workbook()
    ws = wb.active
    ws.title = "HRC NWE weekly"
    ws.append(["Week ending", "HRC NWE (EUR/t)", "Comment"])
    for d in days:
        if d.weekday() == 4:
            ws.append([d.strftime("%d/%m/%Y"), round(idx["HRC"][d], 1), ""])
    ws.append(["", "", "Source: weekly assessment, delayed publication"])
    bold_header(ws)
    wb.save(out / "steel_hrc_weekly.xlsx")

    # PA66: monthly, "Sep-26" style months, EUR/kg with decimal comma, semicolon separated
    months = sorted({(d.year, d.month) for d in days})
    rows = []
    for y, m in months:
        label = date(y, m, 1).strftime("%b-%y")
        rows.append((label, f"{month_avg(idx['PA66'], y, m):.3f}".replace(".", ","), "contract"))
    write_csv(out / "pa66_europe_monthly.csv", ["Month", "PA66 EUR/kg", "Basis"], rows, delimiter=";")

    # Rubber: US-style dates, US cents/kg
    rows = [(d.strftime("%m/%d/%Y"), round(idx["RUBBER"][d], 1)) for d in days]
    write_csv(out / "sgx_smr20_rubber.csv", ["Date", "Settle (USc/kg)"], rows)

    # FX
    rows = [(d.isoformat(), round(idx["EURUSD"][d], 4), round(idx["EURPLN"][d], 4)) for d in days]
    write_csv(out / "fx_reference_rates.csv", ["date", "EURUSD", "EURPLN"], rows)

    # Analyst notes (document-based market intelligence)
    notes = out / "analyst_notes"
    notes.mkdir(parents=True, exist_ok=True)
    alu_jul = month_avg(idx["ALU"], 2026, 7)
    alu_sep = month_avg(idx["ALU"], 2026, 9)
    (notes / "2026-09-08_aluminium_outlook.md").write_text(
        "# Aluminium outlook - September 2026\n\n"
        f"LME aluminium cash averaged about {alu_jul:,.0f} USD/t in July and has since climbed to around "
        f"{alu_sep:,.0f} USD/t month-to-date in September.\n\n"
        "Drivers: smelter curtailments in Europe after an energy price spike, and lower Chinese exports.\n"
        "Our base case is that prices stay elevated through Q4 2026, with upside risk if curtailments extend.\n\n"
        "Implication for buyers: index-linked contracts with monthly review will pass the increase through "
        "from October. Check which customer agreements allow recovery and file claims on time.\n"
    )
    (notes / "2026-09-15_copper_flash.md").write_text(
        "# Flash note - copper\n\n"
        "Copper cash prices rose sharply in the first half of September on mine disruptions in South America.\n"
        "Wire rod premiums in Central Europe are also up. Expect monthly-review contracts to adjust from October.\n"
    )
    (notes / "2026-09-01_polymers_rubber.md").write_text(
        "# Polymers and elastomers - monthly review\n\n"
        "PA66 contract prices increased again in August on tight adiponitrile supply; further increases are "
        "announced for Q4. Natural rubber firmed modestly on weather-related tapping delays.\n"
        "Steel HRC in Northwest Europe drifted lower on weak construction demand.\n"
    )


def write_sap_bw(idx):
    out = RAW / "sap_bw"
    # Material master: SAP column names, semicolon, leading-zero MATNR
    rows = [(m[2], m[1], "KG", m[4], "1000") for m in MATERIALS]
    rows.append(("000000000040001003", "ALU A380 INGOT (OBSOLETE)", "KG", "RM-NF", "1000"))  # injected: obsolete dup
    write_csv(out / "0MATERIAL_master.csv", ["MATNR", "MAKTX", "MEINS", "MATKL", "WERKS"], rows, delimiter=";")

    rows = [(v[0], v[1], v[2], v[3]) for v in VENDORS.values()]
    write_csv(out / "0VENDOR_master.csv", ["LIFNR", "NAME1", "LAND1", "WAERS"], rows, delimiter=";")

    # Purchase order history: NETPR per PEINH units (PEINH=1000 means per 1000 KG), mixed currencies
    rows = []
    po = 4500010000
    issues = []
    month = date(2025, 1, 1)
    while month <= date(2026, 9, 1):
        for vkey, (lifnr, _name, _c, waers, mats) in VENDORS.items():
            for mcode in mats:
                contract = next(c for c in VENDOR_CONTRACTS if c["vendor"] == vkey)
                price_kg = contract_price_kg(contract, mcode, month, idx)
                for _ in range(rng.randint(1, 3)):
                    po += 1
                    qty_kg = round(contract["volume_t"] * 1000 / 12 / 2 * rng.uniform(0.6, 1.4) /
                                   max(1, len(mats)), -1)
                    d = month + timedelta(days=rng.randint(0, 25))
                    peinh = 1000 if mcode.startswith("RM-ST") or rng.random() < 0.3 else 1
                    netpr = price_kg * peinh
                    if waers == "PLN":
                        netpr *= idx["EURPLN"].get(d, 4.3)
                    rows.append([str(po), "10", MAT[mcode][2], lifnr, f"{qty_kg:.0f}", "KG",
                                 f"{netpr:.2f}", str(peinh), waers, d.strftime("%Y%m%d"), "1000"])
        month = (month.replace(day=28) + timedelta(days=4)).replace(day=1)

    # injected issues
    rows[57][6] = f"{float(rows[57][6]) * 10:.2f}"
    issues.append(f"PO {rows[57][0]}: NETPR is 10x too high (decimal slip)")
    rows.append(list(rows[120]))
    issues.append(f"PO {rows[120][0]}: exact duplicate line")
    rows[200][3] = ""
    issues.append(f"PO {rows[200][0]}: LIFNR missing")
    rows[260][5] = "TO"
    rows[260][4] = f"{float(rows[260][4]) / 1000:.1f}"
    issues.append(f"PO {rows[260][0]}: quantity in TO (tonnes) instead of KG")
    write_csv(out / "2LIS_02_ITM_po_history.csv",
              ["EBELN", "EBELP", "MATNR", "LIFNR", "MENGE", "MEINS", "NETPR", "PEINH", "WAERS", "BEDAT", "WERKS"],
              rows, delimiter=";")
    return issues


def contract_price_kg(contract, mcode, month, idx):
    """Price in contract currency per kg for a given delivery month (simplified formula)."""
    base = contract["base_price"]
    if contract["price_unit"].startswith("EUR/t"):
        base = base / 1000 + (0.095 if mcode == "RM-ST-CRC" else 0)
    if not contract["index"]:
        return base
    prev = (month.replace(day=1) - timedelta(days=1))
    ref = contract["base_index_value"]
    if prev < START:
        return base
    cur = month_avg(idx[contract["index"]], prev.year, prev.month)
    change = (cur - ref) / ref
    if abs(change) * 100 <= contract["trigger_pct"]:
        return base
    adj = change * contract["share"]
    if contract["cap_pct"]:
        adj = max(-contract["cap_pct"] / 100, min(contract["cap_pct"] / 100, adj))
    return base * (1 + adj)


def write_purchasing_excel():
    """Buyer-maintained workbook: free-text names, mixed units, conflicts with SAP."""
    out = RAW / "purchasing"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 2026 prices"
    ws.append(["Supplier", "Material", "Price", "Unit", "Valid from", "Buyer", "Notes"])
    rows = [
        ("AluCast Iberia", "AL A380", 2.85, "EUR/kg", "01.07.2026", "J. Weber", "no change since Q1"),
        ("Midwest Diecast", "Alu A380", 3180, "USD/t", "2026-07-01", "S. Patel", "index adj. July"),
        ("Nordic Rolled Alu", "Alu sheet 6061", 4.40, "EUR/kg", "1/7/2026", "J. Weber", ""),
        ("Baltic Steel", "HRC S355", 690, "EUR/t", "01.07.2026", "M. Kask", "amendment A1 signed"),
        ("Baltic Steel", "CRC DC04", 0.785, "EUR/kg", "01.07.2026", "M. Kask", ""),
        ("Cuprum Wire", "Cu rod", 39.9, "PLN/kg", "01.07.2026", "A. Nowak", "invoice PLN"),
        ("PolyNova", "PA66 GF30", 3.95, "EUR/kg", "01.04.2026", "J. Weber", "quarterly review pending"),
        ("Keystone Rubber", "SMR20", 2.10, "USD/kg", "01.01.2026", "S. Patel", "contract ends Dec - renew?"),
        ("Zenith Zinc", "Zinc SHG", None, "EUR/kg", "", "M. Kask", "waiting for supplier"),
        ("AluCast Iberia", "A380 ingot", 2.95, "EUR/kg", "15.08.2026", "J. Weber", "verbal quote, not in SAP"),
    ]
    for r in rows:
        ws.append(list(r))
    bold_header(ws)
    wb.save(out / "buyer_price_list_Q3_2026.xlsx")
    return [
        "buyer_price_list: AluCast A380 has two conflicting rows (2.85 vs 2.95, the latter a verbal quote)",
        "buyer_price_list: Zinc price missing",
        "buyer_price_list: mixed units (EUR/kg, USD/t, PLN/kg) and three date formats",
        "buyer_price_list: PolyNova price is stale (valid from Q2) despite PA66 index rise",
    ]


def write_bom():
    out = RAW / "bom"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM v7"
    ws.append(["Parent", "Level", "Component", "Description", "Qty", "UoM", "Version", "Valid from"])
    desc = {**PRODUCTS, **{m[0]: m[1] for m in MATERIALS},
            "SA-BC-HOUSING": "Caliper housing casting", "SA-BC-BRACKET": "Caliper bracket stamping",
            "SA-BE-TRAY": "Battery tray weldment"}
    for parent, child, qty, uom, lvl in BOM:
        # injected: rubber seal expressed in grams, zinc coating in lb
        if child == "RM-NR-SMR20" and parent == "FG-BC-200":
            qty, uom = 50, "G"
        if child == "RM-ZN-SHG":
            qty, uom = round(0.12 / 0.45359237, 3), "LB"
        ws.append([parent, lvl, child, desc[child], qty, uom, "v7", "2026-01-01"])
    # injected: superseded version row still present
    ws.append(["SA-BE-TRAY", 2, "RM-AL-6061", desc["RM-AL-6061"], 19.5, "KG", "v6", "2025-01-01"])
    bold_header(ws)
    wb.save(out / "bom_export_v7.xlsx")
    return [
        "bom: FG-BC-200 rubber seal in G (50 g) instead of KG",
        "bom: zinc coating qty in LB",
        "bom: superseded v6 row for SA-BE-TRAY (19.5 kg) alongside v7 (18.0 kg)",
    ]


def write_sales():
    out = RAW / "sales"
    months = [date(2026, 10, 1)]
    for _ in range(11):
        m = months[-1]
        months.append((m.replace(day=28) + timedelta(days=4)).replace(day=1))
    seasonal = [1.0, 1.05, 0.8, 0.9, 1.0, 1.1, 1.05, 1.0, 1.0, 0.7, 0.85, 1.1]

    # EU plant: semicolon CSV, DD.MM.YYYY, decimal comma, customer aliases
    rows = []
    for ckey in ("C-NORDWERK", "C-SAKURA"):
        c = CUSTOMERS[ckey]
        for p, annual in c["products"].items():
            for m, s in zip(months, seasonal):
                qty = annual / 12 * s * rng.uniform(0.9, 1.1)
                rows.append((m.strftime("%d.%m.%Y"), rng.choice(c["aliases"]), p, f"{qty:.1f}".replace(".", ",")))
    write_csv(out / "forecast_EU_plant_2026-10.csv", ["Monat", "Kunde", "Material", "Menge_Stk"], rows, delimiter=";")

    # US plant: Excel, MM/DD/YYYY, product codes without FG- prefix
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.append(["Month", "Customer", "Part #", "Units"])
    c = CUSTOMERS["C-BRIGHTLINE"]
    for p, annual in c["products"].items():
        for m, s in zip(months, seasonal):
            ws.append([m.strftime("%m/%d/%Y"), c["aliases"][0], p.replace("FG-", ""),
                       round(annual / 12 * s * rng.uniform(0.9, 1.1))])
    ws.append(["10/01/2026", "Brightline Motors", "BE-500", None])  # injected: blank units
    bold_header(ws)
    wb.save(out / "forecast_US_plant_2026-10.xlsx")
    return [
        "sales: EU file uses 3 customer name variants per customer and decimal commas",
        "sales: US file drops the FG- prefix on part numbers",
        "sales: US file has a row with blank units",
    ]


def write_rmi_tracker():
    out = RAW / "rmi_tracker"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Material Map"
    ws.append(["Tracker name", "BOM code", "SAP MATNR", "Index", "Warn %", "Critical %"])
    thresholds = {"ALU": (5, 10), "CU": (4, 8), "ZN": (5, 10), "HRC": (5, 10), "PA66": (6, 12), "RUBBER": (8, 15)}
    for code, name in TRACKER_NAMES.items():
        m = MAT[code]
        matnr = m[2].lstrip("0")  # injected: leading zeros stripped by Excel
        w, c = thresholds[m[3]]
        ws.append([name, code, int(matnr), m[3], w, c])
    bold_header(ws)

    ws = wb.create_sheet("Overrides")
    ws.append(["Date", "Material", "Override price", "Unit", "Reason", "Approved by"])
    ws.append(["2026-05-12", "Alu A380 ingot", 2.90, "EUR/kg", "Spot buy during AluCast outage", "K. Brandt"])
    ws.append(["2026-08-03", "PA66 GF30", 4.10, "EUR/kg", "Expected Q3 increase (manual)", ""])  # no approver
    ws.append(["2026-08-20", "Copper rod", 9.60, "EUR/kg", "Use for Q4 costing", "K. Brandt"])
    bold_header(ws)

    ws = wb.create_sheet("Monthly RMI")
    ws.append(["Month", "Material", "Index avg", "Base index", "Change %", "Status"])
    for i, (code, name) in enumerate(TRACKER_NAMES.items()):
        r = i + 2
        base = {"ALU": 2350, "CU": 9100, "ZN": 2650, "HRC": 640, "PA66": 3.10, "RUBBER": 170}[MAT[code][3]]
        ws.append(["2026-08", name, None, base, f"=IFERROR((C{r}-D{r})/D{r}*100,\"\")",
                   f"=IF(E{r}=\"\",\"missing\",IF(ABS(E{r})>=VLOOKUP(B{r},'Material Map'!A:F,6,FALSE),"
                   f"\"CRITICAL\",IF(ABS(E{r})>=VLOOKUP(B{r},'Material Map'!A:F,5,FALSE),\"WARN\",\"OK\")))"])
    bold_header(ws)
    wb.save(out / "RMI_Tracker_2026.xlsx")
    return [
        "rmi_tracker: SAP MATNR stored as number (leading zeros lost)",
        "rmi_tracker: PA66 override has no approver",
        "rmi_tracker: 'Monthly RMI' index averages for 2026-08 were never filled in (manual step skipped)",
    ]


def write_excel_inputs():
    """Ad-hoc business Excel input: plant-level scrap/yield factors maintained by controlling."""
    out = RAW / "excel_inputs"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Yield factors"
    ws.append(["Plant", "Material", "Yield %", "Owner", "Last reviewed"])
    for plant in ("Wolfsried", "Greenville"):
        for code in MAT:
            y = round(rng.uniform(88, 98), 1)
            ws.append([plant, TRACKER_NAMES[code], y, "Controlling", "2026-03-31"])
    ws.append(["Greenville", "Alu A380 ingot", 0.93, "Controlling", "2026-06-30"])  # injected: fraction not %
    bold_header(ws)
    wb.save(out / "yield_factors_2026.xlsx")
    return ["excel_inputs: Greenville A380 yield entered as 0.93 (fraction) instead of 93 (%), and duplicated"]


def main():
    days, idx = build_indices()
    write_market(days, idx)
    issues = []
    issues += write_sap_bw(idx)
    issues += write_purchasing_excel()
    issues += write_bom()
    issues += write_sales()
    issues += write_rmi_tracker()
    issues += write_excel_inputs()
    write_vendor_contracts()
    write_customer_contracts()
    write_vendor_notice_email(idx)
    issues += [
        "market: LME copper missing 2026-03-17",
        "market: LME aluminium 2026-06-10 appears twice with conflicting values",
        "market: HRC weekly sheet has a trailing footer row",
        "contracts: VC-2025-007 amended by VC-2025-007-A1 (trigger 5% -> 3%, quarterly -> monthly, from 2026-07-01)",
        "contracts: VC-2024-040 (Keystone rubber) is fixed-price and expires 2026-12-31",
        "contracts: VC-2025-022 (Midwest) expires 2026-12-31",
        f"contracts: AluCast email notice claims {EMAIL['claimed']:.2f} EUR/kg from 2026-10-01; "
        f"contract formula gives {EMAIL['correct']:.2f} EUR/kg (overcharge)",
        "customer: Brightline excludes PA66/steel/rubber; Sakura excludes copper -> unrecoverable exposure",
    ]

    summary = {
        "generated_for": TODAY.isoformat(),
        "monthly_index_averages": {
            name: {f"{y}-{m:02d}": round(month_avg(s, y, m), 4)
                   for (y, m) in sorted({(d.year, d.month) for d in days})}
            for name, s in idx.items()
        },
    }
    (ROOT / "reference").mkdir(parents=True, exist_ok=True)
    (ROOT / "reference" / "index_monthly_averages.json").write_text(json.dumps(summary, indent=2))
    (ROOT / "reference" / "master_data.json").write_text(json.dumps({
        "materials": [dict(zip(["code", "description", "matnr", "index", "matkl"], m)) for m in MATERIALS],
        "products": PRODUCTS,
        "bom": [dict(zip(["parent", "child", "qty", "uom", "level"], b)) for b in BOM],
        "customers": CUSTOMERS,
        "vendors": {k: dict(zip(["lifnr", "name", "country", "currency", "materials"], v))
                    for k, v in VENDORS.items()},
        "vendor_contracts": VENDOR_CONTRACTS,
        "customer_contracts": CUSTOMER_CONTRACTS,
    }, indent=2))

    lines = ["# Ground truth: injected data-quality issues", "",
             "Generated by `scripts/generate_sample_data.py` (seed 42). The agent should find these.", ""]
    lines += [f"- {i}" for i in issues]
    alu = idx["ALU"]
    lines += ["", "## Market shock (the demo scenario)", "",
              f"- ALU July 2026 avg: {month_avg(alu, 2026, 7):,.0f} USD/t; "
              f"Aug: {month_avg(alu, 2026, 8):,.0f}; Sep MTD: {month_avg(alu, 2026, 9):,.0f}",
              f"- CU Aug avg: {month_avg(idx['CU'], 2026, 8):,.0f}; Sep MTD: {month_avg(idx['CU'], 2026, 9):,.0f}",
              f"- PA66 Jul: {month_avg(idx['PA66'], 2026, 7):.3f}; Sep: {month_avg(idx['PA66'], 2026, 9):.3f} EUR/kg",
              f"- HRC Jun: {month_avg(idx['HRC'], 2026, 6):.0f}; Sep: {month_avg(idx['HRC'], 2026, 9):.0f} EUR/t"]
    (ROOT / "GROUND_TRUTH.md").write_text("\n".join(lines) + "\n")
    print("\n".join(lines))


if __name__ == "__main__":
    main()
