"""Buyer-maintained price list: free-text names, mixed units and date formats, cross-checked against SAP."""

from collections import defaultdict
from datetime import datetime, timezone

from openpyxl import load_workbook

from rmi.ingest.common import RAW, Context, parse_date, parse_number, price_to_per_kg, to_dt

PATH = RAW / "purchasing" / "buyer_price_list_Q3_2026.xlsx"
SAP_TOLERANCE = 0.02


def load(ctx: Context, as_of: datetime):
    ws = load_workbook(PATH, data_only=True).active
    quarter_start = datetime(as_of.year, 3 * ((as_of.month - 1) // 3) + 1, 1, tzinfo=timezone.utc)
    formats = set()
    quotes = []
    for i, (supplier, material, price, unit, valid_from, buyer, notes) in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2):
        src = ctx.source(PATH, "purchasing", row=i, sheet=ws.title)
        ctx.count(PATH, 1)
        vid, mid = ctx.resolve_vendor(supplier), ctx.resolve_material(material)
        if not vid or not mid:
            ctx.issue(src, "unmapped_identifier", "error",
                      f"Cannot resolve supplier '{supplier}' or material '{material}'",
                      record={"supplier": supplier, "material": material})
            continue
        value = parse_number(price)
        if value is None:
            ctx.issue(src, "missing_required_field", "error",
                      f"{supplier} / {material}: price is blank (note: '{notes}')",
                      entity={"vendor_id": vid, "material_id": mid})
            continue
        # buyers here work in European conventions, so 1/7/2026 is read as 1 July
        d, fmt, ambiguous = parse_date(valid_from, prefer="uk")
        formats.add(fmt)
        if ambiguous:
            ctx.issue(src, "ambiguous_date", "warning",
                      f"'{valid_from}' could be {d:%d %b %Y} or month/day; read as {d:%d %b %Y} (EU buyer)",
                      auto_fixed=True, entity={"vendor_id": vid, "material_id": mid})
        per_kg, currency = price_to_per_kg(value, unit)
        quotes.append({
            "material_id": mid, "vendor_id": vid, "kind": "buyer_price_list",
            "price": {"per_kg": round(per_kg, 6), "currency": currency,
                      "per_kg_eur": round(ctx.to_eur(per_kg, currency, d), 6),
                      "as_entered": {"value": value, "unit": unit}},
            "valid_from": to_dt(d), "buyer": buyer, "notes": notes,
            "status": "verbal" if notes and "verbal" in notes.lower() else "recorded",
            "source": src,
        })
        ctx.count(PATH, 0, 1)

    if len(formats) > 1:
        ctx.issue(ctx.source(PATH, "purchasing", sheet=ws.title), "inconsistent_format", "info",
                  f"'Valid from' uses {len(formats)} date formats ({', '.join(sorted(formats))}); normalised",
                  auto_fixed=True)

    by_pair = defaultdict(list)
    for q in quotes:
        by_pair[(q["vendor_id"], q["material_id"])].append(q)
    for (vid, mid), qs in by_pair.items():
        if len({q["price"]["per_kg_eur"] for q in qs}) > 1:
            desc = "; ".join(f"{q['price']['as_entered']['value']} {q['price']['as_entered']['unit']} from "
                             f"{q['valid_from']:%Y-%m-%d} ({q['status']})" for q in qs)
            ctx.issue(qs[-1]["source"], "conflicting_values", "error",
                      f"{vid} / {mid} has {len(qs)} different buyer prices: {desc}",
                      entity={"vendor_id": vid, "material_id": mid})
        latest = max(qs, key=lambda q: q["valid_from"])
        if latest["valid_from"] < quarter_start:
            ctx.issue(latest["source"], "stale_value", "warning",
                      f"{vid} / {mid} buyer price last updated {latest['valid_from']:%Y-%m-%d}, "
                      f"before the current quarter ({quarter_start:%Y-%m-%d})",
                      entity={"vendor_id": vid, "material_id": mid})
        sap = latest_po_price(ctx, vid, mid)
        recorded = [q for q in qs if q["status"] == "recorded"]
        if sap and recorded:
            q = max(recorded, key=lambda q: q["valid_from"])
            diff = q["price"]["per_kg_eur"] / sap["price"]["per_kg_eur"] - 1
            if abs(diff) > SAP_TOLERANCE:
                ctx.issue(q["source"], "cross_source_mismatch", "warning",
                          f"{vid} / {mid}: buyer list {q['price']['per_kg_eur']:.3f} EUR/kg vs latest SAP PO "
                          f"{sap['_id']} {sap['price']['per_kg_eur']:.3f} EUR/kg ({diff:+.1%})",
                          entity={"vendor_id": vid, "material_id": mid, "po_line": sap["_id"]})
    ctx.docs["price_quotes"].extend(quotes)


def latest_po_price(ctx, vid, mid):
    lines = [l for l in ctx.docs["po_lines"]
             if l["vendor_id"] == vid and l["material_id"] == mid and l["quality"]["status"] == "ok"]
    return max(lines, key=lambda l: l["order_date"]) if lines else None
