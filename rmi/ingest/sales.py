"""Sales forecasts from the EU plant (CSV) and the US plant (Excel)."""

import csv

from openpyxl import load_workbook

from rmi.ingest.common import RAW, Context, parse_date, parse_number, to_dt

DIR = RAW / "sales"
_prefixed = set()  # log each unprefixed part number once per run


def load(ctx: Context):
    _prefixed.clear()
    path = DIR / "forecast_EU_plant_2026-10.csv"
    with path.open(encoding="utf-8") as f:
        for i, r in enumerate(csv.DictReader(f, delimiter=";"), start=2):
            add(ctx, path, i, r["Monat"], r["Kunde"], r["Material"], r["Menge_Stk"], "Wolfsried", prefer="de")

    path = DIR / "forecast_US_plant_2026-10.xlsx"
    ws = load_workbook(path, data_only=True).active
    for i, (month, customer, part, units) in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        add(ctx, path, i, month, customer, part, units, "Greenville", prefer="us", sheet=ws.title)

    variants = {}
    for f in ctx.docs["forecasts"]:
        variants.setdefault(f["customer_id"], set()).add(f["customer_name_raw"])
    for cid, names in variants.items():
        if len(names) > 1:
            ctx.issue({"system": "sales", "file": "data/raw/sales"}, "identifier_variants", "info",
                      f"{cid} appears under {len(names)} spellings: {sorted(names)}; resolved via crosswalk",
                      auto_fixed=True, entity={"customer_id": cid})


def add(ctx, path, i, month, customer, part, units, plant, prefer, sheet=None):
    src = ctx.source(path, "sales", row=i, sheet=sheet)
    ctx.count(path, 1)
    cid = ctx.resolve_customer(customer)
    if not cid:
        ctx.issue(src, "unmapped_identifier", "error", f"Unknown customer '{customer}'", record={"customer": customer})
        return
    pid = part if part in ctx.docs["products"] else f"FG-{part}"
    if pid not in ctx.docs["products"]:
        ctx.issue(src, "unmapped_identifier", "error", f"Unknown part '{part}'")
        return
    if pid != part and part not in _prefixed:
        _prefixed.add(part)
        ctx.issue(src, "identifier_format", "info", f"Part '{part}' lacks the FG- prefix; resolved to {pid}",
                  auto_fixed=True, entity={"product_id": pid})
    qty = parse_number(units)
    if qty is None:
        ctx.issue(src, "missing_required_field", "error",
                  f"Forecast {customer} / {pid} for {month}: units are blank; row not loaded",
                  entity={"customer_id": cid, "product_id": pid})
        return
    d, _, _ = parse_date(month, prefer=prefer)
    ctx.docs["forecasts"].append({
        "customer_id": cid, "product_id": pid, "plant": plant, "month": to_dt(d.replace(day=1)),
        "units": qty, "customer_name_raw": customer, "source": src,
    })
    ctx.count(path, 0, 1)
