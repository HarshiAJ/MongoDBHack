"""RMI Tracker workbook: material crosswalk, thresholds, manual overrides and the monthly sheet."""

from openpyxl import load_workbook

from rmi.ingest.common import CROSSWALK, RAW, Context, norm, parse_date, price_to_per_kg, to_dt

PATH = RAW / "rmi_tracker" / "RMI_Tracker_2026.xlsx"


def load(ctx: Context):
    wb = load_workbook(PATH)  # formulas, not cached values: the tracker has never been recalculated

    ws = wb["Material Map"]
    restored = {}
    for i, (name, code, matnr, index, warn, crit) in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        src = ctx.source(PATH, "rmi_tracker", row=i, sheet=ws.title)
        ctx.count(PATH, 1, 1)
        matnr_str = str(matnr)
        if isinstance(matnr, (int, float)) or len(matnr_str) != 18:
            matnr_str = matnr_str.split(".")[0].zfill(18)
            restored[code] = matnr_str
        ctx.docs["materials"][code] = {
            "_id": code,
            "index": index,
            "uom": "KG",
            "identifiers": {"sap_matnr": matnr_str, "tracker_name": name,
                            "aliases": CROSSWALK["material_aliases"].get(code, [])},
            "thresholds": {"warn_pct": warn, "critical_pct": crit},
            "yield": [],
            "sources": [src],
        }
        ctx.matnr_to_material[matnr_str] = code
        ctx.material_alias[norm(name)] = code

    if restored:
        ctx.issue(ctx.source(PATH, "rmi_tracker", sheet="Material Map"), "identifier_format", "warning",
                  f"SAP MATNR stored as a number for {len(restored)} materials (Excel dropped the leading zeros); "
                  f"restored to 18 digits", auto_fixed=True, fix={"sap_matnr": restored})

    ws = wb["Overrides"]
    for i, (d, name, price, unit, reason, approver) in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        src = ctx.source(PATH, "rmi_tracker", row=i, sheet=ws.title)
        ctx.count(PATH, 1)
        mid = ctx.resolve_material(name)
        if not mid:
            ctx.issue(src, "unmapped_identifier", "error", f"Override for unknown material '{name}'")
            continue
        per_kg, currency = price_to_per_kg(price, unit)
        valid_from, _, _ = parse_date(d)
        quote = {
            "material_id": mid, "vendor_id": None, "kind": "manual_override",
            "price": {"per_kg": per_kg, "currency": currency, "per_kg_eur": per_kg if currency == "EUR" else None},
            "valid_from": to_dt(valid_from), "reason": reason, "approved_by": approver,
            "status": "approved" if approver else "unapproved", "source": src,
        }
        if not approver:
            ctx.issue(src, "missing_approval", "error",
                      f"Manual override {per_kg} {currency}/kg for {mid} has no approver but feeds costing",
                      entity={"material_id": mid})
        ctx.docs["price_quotes"].append(quote)
        ctx.count(PATH, 0, 1)

    ws = wb["Monthly RMI"]
    empty = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        ctx.count(PATH, 1)
        if row[2] is None:
            empty.append(row[1])
    if empty:
        ctx.issue(ctx.source(PATH, "rmi_tracker", sheet=ws.title), "manual_step_skipped", "warning",
                  f"'Monthly RMI' index averages for {ws['A2'].value} are empty for {len(empty)} materials; "
                  f"status formulas evaluate to 'missing'. The pipeline computes them from market_prices instead.",
                  auto_fixed=True, record={"materials": empty})
