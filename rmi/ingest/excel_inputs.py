"""Controlling's yield-factor workbook, attached to materials per plant."""

from openpyxl import load_workbook

from rmi.ingest.common import RAW, Context, parse_date, to_dt

PATH = RAW / "excel_inputs" / "yield_factors_2026.xlsx"


def load(ctx: Context):
    ws = load_workbook(PATH, data_only=True).active
    seen = {}
    for i, (plant, name, pct, owner, reviewed) in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        src = ctx.source(PATH, "excel_inputs", row=i, sheet=ws.title)
        ctx.count(PATH, 1)
        mid = ctx.resolve_material(name)
        if not mid:
            ctx.issue(src, "unmapped_identifier", "error", f"Unknown material '{name}'")
            continue
        if pct <= 1:
            ctx.issue(src, "unit_mismatch", "warning",
                      f"{plant} / {mid} yield entered as {pct} (fraction); interpreted as {pct * 100:.0f}%",
                      auto_fixed=True, entity={"material_id": mid})
            pct = pct * 100
        d, _, _ = parse_date(reviewed)
        entry = {"plant": plant, "yield_pct": pct, "owner": owner, "reviewed": to_dt(d), "source": src}
        key = (plant, mid)
        if key in seen:
            prev = seen[key]
            newer, older = (entry, prev) if entry["reviewed"] >= prev["reviewed"] else (prev, entry)
            ctx.issue(src, "duplicate_record", "warning",
                      f"{plant} / {mid} has two yield factors ({older['yield_pct']}% reviewed "
                      f"{older['reviewed']:%Y-%m-%d}, {newer['yield_pct']}% reviewed {newer['reviewed']:%Y-%m-%d}); "
                      f"kept the most recent",
                      auto_fixed=True, entity={"material_id": mid})
            seen[key] = newer
        else:
            seen[key] = entry
        ctx.count(PATH, 0, 1)
    for (plant, mid), entry in seen.items():
        ctx.docs["materials"][mid]["yield"].append(entry)
