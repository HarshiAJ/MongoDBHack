"""Multi-level BOM export -> products with an exploded raw-material usage (computed pattern)."""

from collections import defaultdict

from openpyxl import load_workbook

from rmi.ingest.common import RAW, Context, qty_to_kg

PATH = RAW / "bom" / "bom_export_v7.xlsx"


def version_no(v: str) -> int:
    return int(str(v).lstrip("vV") or 0)


def load(ctx: Context):
    ws = load_workbook(PATH, data_only=True).active
    latest = {}  # (parent, component) -> row dict
    for i, (parent, level, comp, desc, qty, uom, version, valid_from) in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2):
        src = ctx.source(PATH, "bom", row=i, sheet=ws.title)
        ctx.count(PATH, 1)
        row = {"parent": parent, "component": comp, "qty": qty, "uom": uom, "version": version,
               "level": level, "source": src}
        key = (parent, comp)
        if key in latest:
            old, new = sorted([latest[key], row], key=lambda r: version_no(r["version"]))
            ctx.issue(old["source"], "superseded_version", "warning",
                      f"BOM {parent} -> {comp}: {old['version']} ({old['qty']} {old['uom']}) superseded by "
                      f"{new['version']} ({new['qty']} {new['uom']}); using {new['version']}",
                      auto_fixed=True, entity={"product_id": parent, "material_id": comp})
            latest[key] = new
        else:
            latest[key] = row

    children = defaultdict(list)
    for (parent, comp), row in latest.items():
        if row["uom"] not in ("KG", "EA"):
            kg = qty_to_kg(row["qty"], row["uom"])
            ctx.issue(row["source"], "unit_conversion", "warning",
                      f"BOM {parent} -> {comp} quantity {row['qty']} {row['uom']} converted to {kg:.4f} KG",
                      auto_fixed=True, entity={"product_id": parent, "material_id": comp})
            row = {**row, "qty": kg, "uom": "KG"}
        children[parent].append(row)
        ctx.count(PATH, 0, 1)

    def explode(node, factor, path, out):
        for row in children.get(node, []):
            if row["uom"] == "EA":
                explode(row["component"], factor * row["qty"], path + [row["component"]], out)
            else:
                mid = row["component"]
                if mid not in ctx.docs["materials"]:
                    ctx.issue(row["source"], "unmapped_identifier", "error", f"BOM component {mid} is not a known material")
                    continue
                out[mid]["kg_per_unit"] += factor * row["qty"]
                out[mid]["paths"].append(" > ".join(path + [mid]))

    for parent in [p for p in children if p.startswith("FG-")]:
        usage = defaultdict(lambda: {"kg_per_unit": 0.0, "paths": []})
        explode(parent, 1, [parent], usage)
        ctx.docs["products"][parent] = {
            "_id": parent,
            "description": None,  # filled from customer contracts (documents loader)
            "bom_version": max(version_no(r["version"]) for r in children[parent]),
            "bom": [{k: r[k] for k in ("parent", "component", "qty", "uom", "version", "level")}
                    for p, rows_ in children.items() if p == parent or p in _subassemblies(children, parent)
                    for r in rows_],
            "materials": [{"material_id": m, "kg_per_unit": round(v["kg_per_unit"], 6), "paths": v["paths"]}
                          for m, v in sorted(usage.items())],
            "source": ctx.source(PATH, "bom", sheet=ws.title),
        }


def _subassemblies(children, root):
    out, stack = set(), [root]
    while stack:
        for r in children.get(stack.pop(), []):
            if r["uom"] == "EA":
                out.add(r["component"])
                stack.append(r["component"])
    return out
